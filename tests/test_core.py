# SPDX-License-Identifier: MIT
"""Regression tests for the scientific core and bundled equation library."""

from __future__ import annotations

import ast
import json
import math
from pathlib import Path
import re
from tempfile import TemporaryDirectory
import unittest

import numpy as np
from openpyxl import load_workbook

from carbon_calculator.calculations import (
    RangeViolation,
    area_normalized_carbon_density,
    calculate_carbon,
    calculate_site_carbon_metrics,
    project_future_carbon,
)
from carbon_calculator.data import (
    RESTORATION_ENVIRONMENTS,
    SHRUB_SPECIES,
    TREE_SPECIES,
    tree_species_for_env,
)
from carbon_calculator.data2 import DOMESTIC_SPECIES, FOREIGN_SPECIES
from carbon_calculator.equation_eval import EvaluationError, evaluate
from carbon_calculator.excel_export import export_all_regions_to_excel
from carbon_calculator.i18n import missing_scientific_names, tr
from carbon_calculator.input_limits import (
    SHRUB_PLANTING_AREA_M2_PER_INDIVIDUAL,
    TREE_PLANTING_AREA_M2_PER_INDIVIDUAL,
    planting_area_budget,
)
from carbon_calculator.tree_simulation.growth_models import render_states
from carbon_calculator.tree_simulation.models import VisualizationInputGroup
from carbon_calculator.tree_simulation.snapshot import build_snapshot
from carbon_calculator.translations import EN
from carbon_calculator.version import __version__


REPOSITORY_ROOT = Path(__file__).resolve().parents[1]


class LibraryTests(unittest.TestCase):
    def test_korean_ui_literals_have_english_translations(self) -> None:
        missing: set[str] = set()
        checked = 0
        for path in (REPOSITORY_ROOT / "carbon_calculator").rglob("*.py"):
            tree = ast.parse(path.read_text(encoding="utf-8"), filename=str(path))
            for node in ast.walk(tree):
                if not isinstance(node, ast.Call) or not node.args:
                    continue
                if not isinstance(node.func, ast.Name) or node.func.id != "tr":
                    continue
                key = node.args[0]
                if not isinstance(key, ast.Constant) or not isinstance(key.value, str):
                    continue
                if not re.search(r"[가-힣]", key.value):
                    continue
                checked += 1
                if key.value not in EN:
                    missing.add(key.value)
        self.assertGreater(checked, 200)
        self.assertEqual(sorted(missing), [])

    def test_all_bundled_species_have_english_scientific_names(self) -> None:
        names = {
            *TREE_SPECIES,
            *SHRUB_SPECIES,
            *DOMESTIC_SPECIES,
            *FOREIGN_SPECIES,
        }
        self.assertEqual(missing_scientific_names(names), [])

    def test_plot_labels_are_translated_at_call_time(self) -> None:
        source = (REPOSITORY_ROOT / "carbon_calculator" / "plotting.py").read_text(
            encoding="utf-8"
        )
        tree = ast.parse(source)
        for node in ast.walk(tree):
            if not isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)):
                continue
            defaults = [*node.args.defaults, *[d for d in node.args.kw_defaults if d]]
            for default in defaults:
                with self.subTest(function=node.name):
                    self.assertFalse(
                        isinstance(default, ast.Call)
                        and isinstance(default.func, ast.Name)
                        and default.func.id == "tr",
                        "Translated defaults must be resolved when the function is called",
                    )

    def test_release_version_is_synchronised(self) -> None:
        cff = (REPOSITORY_ROOT / "CITATION.cff").read_text(encoding="utf-8")
        installer = (REPOSITORY_ROOT / "installer.iss").read_text(encoding="utf-8")
        release_notes = (REPOSITORY_ROOT / "RELEASE_NOTES.md").read_text(encoding="utf-8")
        self.assertIn(f"version: {__version__}", cff)
        self.assertIn(f'#define MyAppVersion "{__version__}"', installer)
        self.assertIn(f"# FORECAST-SW v{__version__}", release_notes)

    def test_release_library_counts(self) -> None:
        self.assertEqual(len(TREE_SPECIES), 7)
        self.assertEqual(len(SHRUB_SPECIES), 15)
        self.assertEqual(len(DOMESTIC_SPECIES), 30)
        self.assertEqual(len(FOREIGN_SPECIES), 25)

    def test_bundled_json_is_parseable_and_licensed(self) -> None:
        payload = json.loads((REPOSITORY_ROOT / "species_data.json").read_text(encoding="utf-8"))
        self.assertEqual(payload["_schema"]["license"], "KOGL Type 1 (Attribution)")
        self.assertEqual(payload["_schema"]["assessment_diameter_unit"], "cm")
        self.assertEqual(payload["_schema"]["TREE_BASE_equation_diameter_unit"], "cm")
        self.assertEqual(payload["_schema"]["SHRUB_SPECIES_equation_diameter_unit"], "mm")
        self.assertEqual(payload["_schema"]["growth_diameter_unit"], "cm/year")
        self.assertEqual(len(payload["TREE_BASE"]), 7)
        self.assertEqual(len(payload["SHRUB_SPECIES"]), 15)

    def test_site_category_does_not_select_coefficients(self) -> None:
        baseline = tree_species_for_env(RESTORATION_ENVIRONMENTS[0])
        for category in RESTORATION_ENVIRONMENTS[1:]:
            self.assertEqual(tree_species_for_env(category), baseline)


class CalculationTests(unittest.TestCase):
    def setUp(self) -> None:
        self.pine = TREE_SPECIES["소나무"]

    def test_allometric_formula_and_exact_count_scaling(self) -> None:
        one = calculate_carbon("소나무", self.pine, 5.0, 1)
        hundred = calculate_carbon("소나무", self.pine, 5.0, 100)
        expected = self.pine.a * 5.0 ** self.pine.b * self.pine.cf
        self.assertIsNotNone(one)
        self.assertIsNotNone(hundred)
        self.assertAlmostEqual(one.carbon_kg, expected, places=12)
        self.assertAlmostEqual(hundred.carbon_kg, one.carbon_kg * 100, places=12)

    def test_area_normalization_matches_the_controlled_site_comparison(self) -> None:
        tree = 194.146028603107
        shrub = 4.743959876150
        total = tree + shrub
        areas = (400.0, 500.0, 600.0)
        metrics = [calculate_site_carbon_metrics(tree, shrub, area) for area in areas]
        densities = [item.area_normalized_kg_m2 for item in metrics]

        self.assertEqual([round(value, 4) for value in densities], [0.4972, 0.3978, 0.3315])
        self.assertEqual([item.total_carbon_kg for item in metrics], [total] * 3)
        for value, area in zip(densities, areas):
            self.assertAlmostEqual(value * area, total, places=12)

        for invalid_area in (0.0, -1.0, math.inf, math.nan):
            with self.subTest(site_area_m2=invalid_area), self.assertRaises(ValueError):
                area_normalized_carbon_density(total, invalid_area)
        with self.assertRaises(ValueError):
            area_normalized_carbon_density(math.nan, 400.0)

    def test_combined_xlsx_exports_area_normalized_density(self) -> None:
        total = 198.889988479257
        comparison_data = [
            {
                "name": f"Profile {index}",
                "area": area,
                "env": RESTORATION_ENVIRONMENTS[0],
                "tree": 194.146028603107,
                "shrub": 4.743959876150,
                "total": total,
                "density": area_normalized_carbon_density(total, area),
            }
            for index, area in enumerate((400.0, 500.0, 600.0), start=1)
        ]

        with TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "controlled_site_comparison.xlsx"
            export_all_regions_to_excel(str(path), [], comparison_data)
            workbook = load_workbook(path, data_only=True)
            sheet = workbook[tr("지역_비교분석")]
            self.assertIn("normalized", str(sheet.cell(2, 7).value).lower())
            self.assertEqual(
                [sheet.cell(row, 6).value for row in range(3, 6)],
                [round(total, 2)] * 3,
            )
            self.assertEqual(
                [sheet.cell(row, 7).value for row in range(3, 6)],
                [0.4972, 0.3978, 0.3315],
            )
            workbook.close()

    def test_zero_quantity_is_excluded(self) -> None:
        self.assertIsNone(calculate_carbon("소나무", self.pine, 5.0, 0))

    def test_diameter_boundaries_are_enforced(self) -> None:
        calculate_carbon("소나무", self.pine, self.pine.diameter_min, 1)
        calculate_carbon("소나무", self.pine, self.pine.diameter_max, 1)
        with self.assertRaises(RangeViolation):
            calculate_carbon("소나무", self.pine, self.pine.diameter_min - 0.1, 1)
        with self.assertRaises(RangeViolation):
            calculate_carbon("소나무", self.pine, self.pine.diameter_max + 0.1, 1)

    def test_projection_is_deterministic_and_matches_year_zero(self) -> None:
        current = calculate_carbon("소나무", self.pine, 5.0, 100)
        years_a, carbon_a = project_future_carbon(self.pine, 5.0, 100, years=50)
        years_b, carbon_b = project_future_carbon(self.pine, 5.0, 100, years=50)
        np.testing.assert_array_equal(years_a, years_b)
        np.testing.assert_array_equal(carbon_a, carbon_b)
        self.assertAlmostEqual(carbon_a[0], current.carbon_kg, places=12)

    def test_every_legacy_shrub_equation_accepts_rcd_in_cm_without_numerical_change(self) -> None:
        payload = json.loads((REPOSITORY_ROOT / "species_data.json").read_text(encoding="utf-8"))
        for name, raw in payload["SHRUB_SPECIES"].items():
            shrub = SHRUB_SPECIES[name]
            a, b, cf, diameter_min_mm, diameter_max_mm, *_growth = raw
            for diameter_mm in (
                diameter_min_mm,
                (diameter_min_mm + diameter_max_mm) / 2.0,
                diameter_max_mm,
            ):
                with self.subTest(species=name, diameter_mm=diameter_mm):
                    diameter_cm = diameter_mm / 10.0
                    row = calculate_carbon(name, shrub, diameter_cm, 3)
                    expected = a * diameter_mm ** b * cf * 3
                    self.assertIsNotNone(row)
                    self.assertAlmostEqual(row.carbon_kg, expected, places=12)
                    self.assertAlmostEqual(row.diameter, diameter_cm, places=12)

    def test_shrub_growth_uses_the_common_centimetre_timeline(self) -> None:
        shrub = SHRUB_SPECIES["사철나무"]
        _, carbon = project_future_carbon(shrub, 1.0, 1, years=1)
        expected_diameter_mm = (1.0 + shrub.growth_y10) * 10.0
        expected = shrub.a * expected_diameter_mm ** shrub.b * shrub.cf
        self.assertAlmostEqual(carbon[1], expected, places=12)

    def test_fractional_shrub_rcd_boundaries_are_enforced_in_cm(self) -> None:
        shrub = SHRUB_SPECIES["병꽃나무"]
        self.assertAlmostEqual(shrub.diameter_min, 0.6)
        self.assertAlmostEqual(shrub.diameter_max, 3.9)
        calculate_carbon("병꽃나무", shrub, 0.6, 1)
        calculate_carbon("병꽃나무", shrub, 3.9, 1)
        with self.assertRaises(RangeViolation):
            calculate_carbon("병꽃나무", shrub, 0.59, 1)
        with self.assertRaises(RangeViolation):
            calculate_carbon("병꽃나무", shrub, 3.91, 1)

    def test_shrub_visualization_uses_centimetres_for_geometry_and_carbon(self) -> None:
        shrub = SHRUB_SPECIES["병꽃나무"]
        item = VisualizationInputGroup("병꽃나무", "shrub", 1.5, 1, "cm", shrub)
        snapshot = build_snapshot(
            region_name="unit-test",
            environment="",
            area_w=5,
            area_h=5,
            inputs=(item,),
        )
        state = render_states(snapshot, 0)[0]
        current = calculate_carbon("병꽃나무", shrub, 1.5, 1)
        self.assertAlmostEqual(snapshot.groups[0].diameter_by_year[0], 1.5)
        self.assertAlmostEqual(
            snapshot.groups[0].diameter_by_year[1], 1.5 + shrub.growth_y10
        )
        self.assertAlmostEqual(snapshot.groups[0].carbon_by_year_kgc[0], current.carbon_kg)
        self.assertAlmostEqual(state.diameter_m, 0.015)

    def test_combined_planting_area_guard_is_enforced_at_the_boundary(self) -> None:
        self.assertEqual(TREE_PLANTING_AREA_M2_PER_INDIVIDUAL, 1.0)
        self.assertEqual(SHRUB_PLANTING_AREA_M2_PER_INDIVIDUAL, 0.25)

        at_limit = planting_area_budget(
            area_w=10, area_h=10, tree_quantity=50, shrub_quantity=200
        )
        self.assertIsNotNone(at_limit)
        self.assertAlmostEqual(at_limit.required_area_m2, 100.0)
        self.assertFalse(at_limit.is_exceeded)
        self.assertAlmostEqual(at_limit.excess_area_m2, 0.0)

        over_limit = planting_area_budget(
            area_w=10, area_h=10, tree_quantity=80, shrub_quantity=100
        )
        self.assertIsNotNone(over_limit)
        self.assertAlmostEqual(over_limit.tree_area_m2, 80.0)
        self.assertAlmostEqual(over_limit.shrub_area_m2, 25.0)
        self.assertAlmostEqual(over_limit.required_area_m2, 105.0)
        self.assertTrue(over_limit.is_exceeded)
        self.assertAlmostEqual(over_limit.excess_area_m2, 5.0)

        self.assertIsNone(
            planting_area_budget(
                area_w=0, area_h=10, tree_quantity=1, shrub_quantity=1
            )
        )


class EquationEvaluatorTests(unittest.TestCase):
    def test_supported_equation_forms(self) -> None:
        self.assertAlmostEqual(evaluate("Y=0.063*X^2.578", 10), 0.063 * 10 ** 2.578)
        self.assertAlmostEqual(
            evaluate("ln(Y)=2.43*ln(X)-2.28", 10),
            math.exp(2.43 * math.log(10) - 2.28),
        )
        self.assertAlmostEqual(
            evaluate("Y=exp(-4.7483+1.7395*ln(X*H))", 10, 12),
            math.exp(-4.7483 + 1.7395 * math.log(120)),
        )

    def test_entire_compatibility_library_executes(self) -> None:
        for name, record in {**DOMESTIC_SPECIES, **FOREIGN_SPECIES}.items():
            if record.has_range:
                x = (record.diameter_min + record.diameter_max) / 2
            else:
                x = 10.0
            h = record.var2_default if record.is_multivar else None
            with self.subTest(name=name):
                self.assertTrue(math.isfinite(evaluate(record.equation, x, h)))

    def test_unsafe_or_unknown_syntax_is_rejected(self) -> None:
        unsafe = (
            "Y=__import__('os').system('echo unsafe')",
            "Y=(1).__class__",
            "Y=[X][0]",
            "Y=UNKNOWN(X)",
        )
        for equation in unsafe:
            with self.subTest(equation=equation), self.assertRaises(EvaluationError):
                evaluate(equation, 10)


if __name__ == "__main__":
    unittest.main()
